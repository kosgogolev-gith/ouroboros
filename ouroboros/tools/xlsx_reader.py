"""XLSX reader tool: read Excel files and return structured text representation."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import List

from ouroboros.tools.registry import ToolContext, ToolEntry

log = logging.getLogger(__name__)


def _resolve_path(ctx: ToolContext, path: str) -> Path:
    """Resolve path: absolute or relative to DRIVE_ROOT."""
    p = Path(path)
    if p.is_absolute():
        return p
    return ctx.drive_path(path)


def _sheet_to_text(ws) -> str:
    """Convert a worksheet to a text table."""
    lines: list[str] = []
    for row in ws.iter_rows(values_only=True):
        cells = []
        for cell in row:
            if cell is None:
                cells.append("")
            else:
                cells.append(str(cell).strip())
        # Skip fully empty rows
        if not any(cells):
            continue
        lines.append(" | ".join(cells))
    return "\n".join(lines)


def _xlsx_read(ctx: ToolContext, path: str, sheet: str = "") -> str:
    """Read XLSX file and return structured text representation.

    Args:
        path: Absolute path or path relative to DRIVE_ROOT.
        sheet: Sheet name to read. If empty, reads all sheets.

    Returns:
        Structured text with sheet contents in table format.
    """
    try:
        import openpyxl
    except ImportError:
        return "\u26a0\ufe0f openpyxl is not installed. Run: pip install openpyxl"

    file_path = _resolve_path(ctx, path)
    if not file_path.exists():
        return f"\u26a0\ufe0f File not found: {file_path}"
    if not file_path.suffix.lower() in (".xlsx", ".xls"):
        return f"\u26a0\ufe0f Not an Excel file: {file_path.name}"

    try:
        wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
    except Exception as e:
        return f"\u26a0\ufe0f Failed to open XLSX: {e}"

    try:
        if sheet:
            if sheet not in wb.sheetnames:
                return (
                    f"\u26a0\ufe0f Sheet '{sheet}' not found. "
                    f"Available sheets: {', '.join(wb.sheetnames)}"
                )
            ws = wb[sheet]
            text = _sheet_to_text(ws)
            return f"=== Sheet: {sheet} ===\n{text}"

        # Read all sheets
        parts: list[str] = []
        for name in wb.sheetnames:
            ws = wb[name]
            text = _sheet_to_text(ws)
            if text:
                parts.append(f"=== Sheet: {name} ===\n{text}")
        if not parts:
            return "\u26a0\ufe0f Workbook is empty (no data in any sheet)."
        return "\n\n".join(parts)
    finally:
        wb.close()


def get_tools() -> List[ToolEntry]:
    return [
        ToolEntry("xlsx_read", {
            "name": "xlsx_read",
            "description": (
                "Read an XLSX (Excel) file and return structured text. "
                "Use for technical specifications, BOMs, and commercial offers. "
                "Supports reading specific sheets or all sheets at once."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {
                        "type": "string",
                        "description": (
                            "Path to XLSX file. Absolute or relative to Drive root."
                        ),
                    },
                    "sheet": {
                        "type": "string",
                        "description": (
                            "Sheet name to read. Leave empty to read all sheets."
                        ),
                    },
                },
                "required": ["path"],
            },
        }, _xlsx_read),
    ]
