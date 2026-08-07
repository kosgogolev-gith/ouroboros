"""Tools for reading spreadsheets and storing or comparing specifications."""

from __future__ import annotations

import difflib
import sqlite3
from pathlib import Path
from typing import Any, List, Sequence

from ouroboros.tools.registry import ToolContext, ToolEntry

try:
    import openpyxl
except ImportError:  # pragma: no cover - depends on optional installation
    openpyxl = None


def _markdown_cell(value: Any) -> str:
    """Make a value safe for a Markdown table cell."""
    if value is None:
        return ""
    return str(value).replace("\\", "\\\\").replace("|", "\\|").replace("\r", "").replace("\n", "<br>")


def _markdown_table(headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> str:
    """Render headers and rows as a Markdown table."""
    header_line = "| " + " | ".join(_markdown_cell(header) for header in headers) + " |"
    separator = "| " + " | ".join("---" for _ in headers) + " |"
    body = ["| " + " | ".join(_markdown_cell(value) for value in row) + " |" for row in rows]
    return "\n".join([header_line, separator, *body])


def _column_name(index: int) -> str:
    """Return an Excel-style column name for a one-based index."""
    name = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        name = chr(65 + remainder) + name
    return name


def _resolve_xlsx_path(ctx: ToolContext, file_path: str) -> Path:
    """Resolve an absolute path or a safe path relative to the repository or Drive."""
    candidate = Path(file_path).expanduser()
    if candidate.is_absolute():
        return candidate.resolve()
    repo_candidate = ctx.repo_path(file_path)
    return repo_candidate if repo_candidate.exists() else ctx.drive_path(file_path)


def _xlsx_read(ctx: ToolContext, file_path: str) -> str:
    """Read the first 50 rows of the active worksheet as a Markdown table."""
    if openpyxl is None:
        return "⚠️ xlsx_read is unavailable because optional dependency 'openpyxl' is not installed."

    try:
        path = _resolve_xlsx_path(ctx, file_path)
    except ValueError as exc:
        return f"⚠️ Invalid file path: {exc}"

    if not path.is_file():
        return f"⚠️ XLSX file not found: {path}"
    if path.suffix.lower() != ".xlsx":
        return "⚠️ xlsx_read accepts only .xlsx files."

    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            worksheet = workbook.active
            source_rows = list(
                worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, 50), values_only=True)
            )
        finally:
            workbook.close()
    except Exception as exc:
        return f"⚠️ Could not read XLSX file: {exc}"

    if not source_rows:
        return f"Workbook '{path.name}', sheet '{worksheet.title}' is empty."

    width = max((len(row) for row in source_rows), default=0)
    while width and all((row[width - 1] if len(row) >= width else None) is None for row in source_rows):
        width -= 1
    if width == 0:
        width = 1

    headers = ["Row", *(_column_name(index) for index in range(1, width + 1))]
    rows = [
        [row_number, *(row[index] if index < len(row) else None for index in range(width))]
        for row_number, row in enumerate(source_rows, start=1)
    ]
    return (
        f"## {path.name} — {worksheet.title}\n\n"
        f"Showing {len(rows)} of {worksheet.max_row} worksheet rows.\n\n"
        f"{_markdown_table(headers, rows)}"
    )


def _spec_compare(ctx: ToolContext, spec1_text: str, spec2_text: str, criteria: str = "") -> str:
    """Compare two specification texts line by line."""
    del ctx
    first_lines = spec1_text.splitlines()
    second_lines = spec2_text.splitlines()
    matcher = difflib.SequenceMatcher(a=first_lines, b=second_lines, autojunk=False)

    differences = []
    for tag, first_start, first_end, second_start, second_end in matcher.get_opcodes():
        if tag == "equal":
            continue
        first_block = first_lines[first_start:first_end]
        second_block = second_lines[second_start:second_end]
        length = max(len(first_block), len(second_block))
        for offset in range(length):
            left = first_block[offset] if offset < len(first_block) else ""
            right = second_block[offset] if offset < len(second_block) else ""
            if tag == "replace":
                kind = "changed"
            elif tag == "delete":
                kind = "only in spec 1"
            else:
                kind = "only in spec 2"
            differences.append((len(differences) + 1, kind, left, right))

    basis = criteria.strip() or "Line-by-line textual comparison"
    summary = (
        f"## Specification comparison\n\n"
        f"**Criteria:** {basis}\n\n"
        f"- Spec 1: {len(first_lines)} lines\n"
        f"- Spec 2: {len(second_lines)} lines\n"
        f"- Differences: {len(differences)}"
    )
    if not differences:
        return summary + "\n\nNo line-level differences found."

    return summary + "\n\n" + _markdown_table(
        ("#", "Difference", "Spec 1", "Spec 2"),
        differences,
    )


def _requirements_db_path(ctx: ToolContext) -> Path:
    """Return the requirements database location, creating its parent directory."""
    directory = ctx.drive_root / "spec_tools"
    directory.mkdir(parents=True, exist_ok=True)
    return directory / "requirements.db"


def _connect_requirements(ctx: ToolContext) -> sqlite3.Connection:
    """Open and initialize the requirements version database."""
    connection = sqlite3.connect(str(_requirements_db_path(ctx)), timeout=10.0)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA journal_mode=WAL")
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS requirements (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_name TEXT NOT NULL,
            requirements_text TEXT NOT NULL,
            created_at TEXT DEFAULT (datetime('now'))
        )
        """
    )
    connection.execute(
        "CREATE INDEX IF NOT EXISTS requirements_project_created ON requirements(project_name, created_at DESC, id DESC)"
    )
    return connection


def _requirements_save(ctx: ToolContext, project_name: str, requirements_text: str) -> str:
    """Save a new requirements version for a project."""
    project_name = project_name.strip()
    if not project_name:
        return "⚠️ Project name must not be empty."

    with _connect_requirements(ctx) as connection:
        cursor = connection.execute(
            "INSERT INTO requirements (project_name, requirements_text) VALUES (?, ?)",
            (project_name, requirements_text),
        )
        version_id = cursor.lastrowid

    return f"✅ Saved requirements for '{project_name}' as version {version_id}."


def _requirements_load(ctx: ToolContext, project_name: str) -> str:
    """Load the newest saved requirements version for a project."""
    project_name = project_name.strip()
    if not project_name:
        return "⚠️ Project name must not be empty."

    with _connect_requirements(ctx) as connection:
        row = connection.execute(
            """
            SELECT id, project_name, requirements_text, created_at
            FROM requirements
            WHERE project_name = ?
            ORDER BY created_at DESC, id DESC
            LIMIT 1
            """,
            (project_name,),
        ).fetchone()

    if row is None:
        return f"No saved requirements found for '{project_name}'."
    return (
        f"## Requirements: {row['project_name']}\n\n"
        f"**Version:** {row['id']}  \n"
        f"**Saved:** {row['created_at']}\n\n"
        f"{row['requirements_text']}"
    )


def get_tools() -> List[ToolEntry]:
    """Expose specification and spreadsheet tools to the registry."""
    return [
        ToolEntry(
            "xlsx_read",
            {
                "name": "xlsx_read",
                "description": "Read the first 50 rows of an .xlsx workbook's active sheet as a Markdown table.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "file_path": {
                            "type": "string",
                            "description": "Absolute path, or path relative to the repository or Drive, of an .xlsx file",
                        }
                    },
                    "required": ["file_path"],
                },
            },
            _xlsx_read,
        ),
        ToolEntry(
            "spec_compare",
            {
                "name": "spec_compare",
                "description": "Compare two specification texts and return a Markdown table of line-level differences.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "spec1_text": {"type": "string", "description": "First specification text"},
                        "spec2_text": {"type": "string", "description": "Second specification text"},
                        "criteria": {
                            "type": "string",
                            "description": "Optional comparison criteria or focus area",
                            "default": "",
                        },
                    },
                    "required": ["spec1_text", "spec2_text"],
                },
            },
            _spec_compare,
        ),
        ToolEntry(
            "requirements_save",
            {
                "name": "requirements_save",
                "description": "Save a new version of project requirements in persistent SQLite storage.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "project_name": {"type": "string", "description": "Project identifier or name"},
                        "requirements_text": {"type": "string", "description": "Requirements text to save"},
                    },
                    "required": ["project_name", "requirements_text"],
                },
            },
            _requirements_save,
        ),
        ToolEntry(
            "requirements_load",
            {
                "name": "requirements_load",
                "description": "Load the newest saved requirements version for a project.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "project_name": {"type": "string", "description": "Project identifier or name"}
                    },
                    "required": ["project_name"],
                },
            },
            _requirements_load,
        ),
    ]
